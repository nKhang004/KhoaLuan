import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, joinedload
from app.core.database import get_db
from app.core.security import get_current_user, require_teacher
from app.models.models import DeThi, CauHoiDeThi, CauHoi, LuaChon, LanThi, NguoiDung

router = APIRouter(prefix="/export", tags=["Xuất dữ liệu"])

HEADER_FILL = PatternFill("solid", fgColor="1B3A6B")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Times New Roman", size=11)
NORMAL_FONT = Font(name="Times New Roman", size=11)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
THIN   = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin")
)

def style_header(cell, text, width=None):
    cell.value = text
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = CENTER
    cell.border = THIN

def style_cell(cell, value="", align=LEFT):
    cell.value = value
    cell.font = NORMAL_FONT
    cell.alignment = align
    cell.border = THIN


@router.get("/exam/{exam_id}/excel", summary="Xuất đề thi ra Excel")
def xuat_de_thi_excel(exam_id: int, db: Session = Depends(get_db), _=Depends(require_teacher)):
    de_thi = db.query(DeThi).filter(DeThi.id == exam_id).first()
    if not de_thi:
        raise HTTPException(status_code=404, detail="Không tìm thấy đề thi")

    cau_hoi_de_thi = (
        db.query(CauHoiDeThi)
        .filter(CauHoiDeThi.de_thi_id == exam_id)
        .order_by(CauHoiDeThi.thu_tu)
        .all()
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Đề thi"
    ws.sheet_view.showGridLines = False

    # Tiêu đề đề thi
    ws.merge_cells("A1:G1")
    title = ws["A1"]
    title.value = f"ĐỀ THI: {de_thi.tieu_de.upper()}"
    title.font = Font(bold=True, size=14, name="Times New Roman", color="1B3A6B")
    title.alignment = CENTER

    ws.merge_cells("A2:G2")
    info = ws["A2"]
    info.value = f"Thời gian: {de_thi.thoi_gian_lam_bai} phút  |  Điểm đạt: {de_thi.diem_dat}/10  |  Tổng số câu: {len(cau_hoi_de_thi)}"
    info.font = Font(italic=True, size=11, name="Times New Roman")
    info.alignment = CENTER
    ws.row_dimensions[2].height = 20

    # Header bảng
    headers = ["STT", "Nội dung câu hỏi", "Độ khó", "Đáp án A", "Đáp án B", "Đáp án C", "Đáp án D"]
    col_widths = [6, 45, 12, 20, 20, 20, 20]
    for col, (h, w) in enumerate(zip(headers, col_widths), 1):
        style_header(ws.cell(row=3, column=col), h)
        ws.column_dimensions[ws.cell(row=3, column=col).column_letter].width = w
    ws.row_dimensions[3].height = 25

    do_kho_map = {"de": "Dễ", "trung_binh": "Trung bình", "kho": "Khó"}
    letters = ["A", "B", "C", "D", "E", "F"]

    for idx, item in enumerate(cau_hoi_de_thi, 1):
        cq = db.query(CauHoi).options(joinedload(CauHoi.lua_chon)).filter(CauHoi.id == item.cau_hoi_id).first()
        if not cq:
            continue
        row = idx + 3
        lua_chon_sorted = sorted(cq.lua_chon, key=lambda x: x.thu_tu)

        style_cell(ws.cell(row=row, column=1), idx, CENTER)
        style_cell(ws.cell(row=row, column=2), cq.noi_dung)
        style_cell(ws.cell(row=row, column=3), do_kho_map.get(cq.do_kho, ""), CENTER)

        for i, lc in enumerate(lua_chon_sorted[:4]):
            marker = " ✓" if lc.la_dap_an else ""
            style_cell(ws.cell(row=row, column=4 + i), f"{letters[i]}. {lc.noi_dung}{marker}")

        ws.row_dimensions[row].height = 35

    # Sheet đáp án
    ws2 = wb.create_sheet("Đáp án")
    ws2.sheet_view.showGridLines = False
    ws2.merge_cells("A1:D1")
    ws2["A1"].value = f"ĐÁP ÁN: {de_thi.tieu_de.upper()}"
    ws2["A1"].font = Font(bold=True, size=13, name="Times New Roman", color="1B3A6B")
    ws2["A1"].alignment = CENTER

    ans_headers = ["STT", "Nội dung câu hỏi", "Đáp án đúng", "Giải thích"]
    ans_widths  = [6, 50, 15, 40]
    for col, (h, w) in enumerate(zip(ans_headers, ans_widths), 1):
        style_header(ws2.cell(row=2, column=col), h)
        ws2.column_dimensions[ws2.cell(row=2, column=col).column_letter].width = w

    for idx, item in enumerate(cau_hoi_de_thi, 1):
        cq = db.query(CauHoi).options(joinedload(CauHoi.lua_chon)).filter(CauHoi.id == item.cau_hoi_id).first()
        if not cq:
            continue
        row = idx + 2
        sorted_lc = sorted(cq.lua_chon, key=lambda x: x.thu_tu)
        dap_an = next((f"{letters[i]}. {lc.noi_dung}" for i, lc in enumerate(sorted_lc[:4]) if lc.la_dap_an), "")
        style_cell(ws2.cell(row=row, column=1), idx, CENTER)
        style_cell(ws2.cell(row=row, column=2), cq.noi_dung)
        style_cell(ws2.cell(row=row, column=3), dap_an, CENTER)
        style_cell(ws2.cell(row=row, column=4), cq.giai_thich or "")
        ws2.row_dimensions[row].height = 30

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"DeThi_{de_thi.id}_{de_thi.tieu_de[:20].replace(' ', '_')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"}
    )


@router.get("/exam/{exam_id}/results/excel", summary="Xuất kết quả thi ra Excel")
def xuat_ket_qua_excel(exam_id: int, db: Session = Depends(get_db), _=Depends(require_teacher)):
    de_thi = db.query(DeThi).filter(DeThi.id == exam_id).first()
    if not de_thi:
        raise HTTPException(status_code=404, detail="Không tìm thấy đề thi")

    lan_thi_list = (
        db.query(LanThi)
        .filter(LanThi.de_thi_id == exam_id, LanThi.nop_bai_luc.isnot(None))
        .order_by(LanThi.tong_diem.desc())
        .all()
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kết quả thi"
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:G1")
    ws["A1"].value = f"KẾT QUẢ THI: {de_thi.tieu_de.upper()}"
    ws["A1"].font = Font(bold=True, size=13, name="Times New Roman", color="1B3A6B")
    ws["A1"].alignment = CENTER

    total = len(lan_thi_list)
    so_dat = sum(1 for lt in lan_thi_list if lt.dat_yeu_cau)
    diem_tb = round(sum(lt.tong_diem for lt in lan_thi_list if lt.tong_diem) / total, 2) if total > 0 else 0

    ws.merge_cells("A2:G2")
    ws["A2"].value = f"Tổng lượt thi: {total}  |  Số đạt: {so_dat}  |  Tỷ lệ đạt: {round(so_dat/total*100,1) if total else 0}%  |  Điểm TB: {diem_tb}"
    ws["A2"].font = Font(italic=True, size=11, name="Times New Roman")
    ws["A2"].alignment = CENTER

    headers  = ["STT", "Tên đăng nhập", "Họ và tên", "Điểm số", "Kết quả", "Bắt đầu", "Nộp bài"]
    widths   = [6, 18, 25, 12, 12, 22, 22]
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        style_header(ws.cell(row=3, column=col), h)
        ws.column_dimensions[ws.cell(row=3, column=col).column_letter].width = w

    PASS_FILL = PatternFill("solid", fgColor="D5F5E3")
    FAIL_FILL = PatternFill("solid", fgColor="FDECEA")

    for idx, lt in enumerate(lan_thi_list, 1):
        hoc_sinh = db.query(NguoiDung).filter(NguoiDung.id == lt.hoc_sinh_id).first()
        row = idx + 3
        fill = PASS_FILL if lt.dat_yeu_cau else FAIL_FILL

        for col in range(1, 8):
            ws.cell(row=row, column=col).fill = fill
            ws.cell(row=row, column=col).border = THIN
            ws.cell(row=row, column=col).font = NORMAL_FONT

        ws.cell(row=row, column=1).value = idx
        ws.cell(row=row, column=1).alignment = CENTER
        ws.cell(row=row, column=2).value = hoc_sinh.ten_dang_nhap if hoc_sinh else "?"
        ws.cell(row=row, column=2).alignment = LEFT
        ws.cell(row=row, column=3).value = hoc_sinh.ho_ten if hoc_sinh else "?"
        ws.cell(row=row, column=3).alignment = LEFT
        ws.cell(row=row, column=4).value = lt.tong_diem
        ws.cell(row=row, column=4).alignment = CENTER
        ws.cell(row=row, column=5).value = "Đạt" if lt.dat_yeu_cau else "Chưa đạt"
        ws.cell(row=row, column=5).alignment = CENTER
        ws.cell(row=row, column=6).value = lt.bat_dau_luc.strftime("%d/%m/%Y %H:%M") if lt.bat_dau_luc else ""
        ws.cell(row=row, column=6).alignment = CENTER
        ws.cell(row=row, column=7).value = lt.nop_bai_luc.strftime("%d/%m/%Y %H:%M") if lt.nop_bai_luc else ""
        ws.cell(row=row, column=7).alignment = CENTER
        ws.row_dimensions[row].height = 22

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"KetQua_{de_thi.id}_{de_thi.tieu_de[:20].replace(' ', '_')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"}
    )
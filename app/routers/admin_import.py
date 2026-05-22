import io
import re
from datetime import datetime, timedelta
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.core.database import get_db
from app.core.security import require_admin, hash_password
from app.models.models import NguoiDung

router = APIRouter(prefix="/admin", tags=["Admin"])

HEADER_FILL = PatternFill("solid", fgColor="1B3A6B")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Times New Roman", size=11)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
THIN = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin")
)

def parse_date(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, (int, float)):
        try:
            excel_base = datetime(1899, 12, 30)
            dt = excel_base + timedelta(days=value)
            return dt.strftime("%d/%m/%Y")
        except:
            pass
    date_str = str(value).strip()
    for sep in ['-', '.']:
        date_str = date_str.replace(sep, '/')
    if re.match(r"^\d{2}/\d{2}/\d{4}$", date_str):
        return date_str
    return None

def find_column_index(row, target_names):
    if not row:
        return None
    for idx, cell in enumerate(row):
        if cell and cell.value:
            cell_str = str(cell.value).strip().lower()
            for target in target_names:
                if target.lower() in cell_str:
                    return idx
    return None

@router.get("/import-template", summary="Tải file Excel mẫu để tạo tài khoản sinh viên")
def tai_file_mau():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Danh sách sinh viên"
    headers = ["Mã sinh viên", "Họ và tên", "Lớp", "Ngày sinh (dd/mm/yyyy)"]
    widths = [15, 35, 20, 15]
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=col)
        cell.value = h
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN
        ws.column_dimensions[cell.column_letter].width = w
    sample = ["SV001", "Nguyễn Văn A", "CNTT-K20", "15/03/2004"]
    for col, val in enumerate(sample, 1):
        cell = ws.cell(row=2, column=col)
        cell.value = val
        cell.font = Font(name="Times New Roman", size=11)
        cell.alignment = LEFT
        cell.border = THIN
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=mau_tao_tai_khoan.xlsx"}
    )

@router.post("/import-users", summary="Import file Excel tạo tài khoản hàng loạt (chỉ sinh viên)")
async def import_users(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user=Depends(require_admin)
):
    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="Chỉ chấp nhận file .xlsx hoặc .xls")
    content = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content))
        ws = wb.active
    except Exception:
        raise HTTPException(status_code=400, detail="Không đọc được file Excel")
    
    header_row = list(ws.iter_rows(min_row=1, max_row=1, values_only=False))[0]
    col_ma = find_column_index(header_row, ["mã sinh viên", "mã sv", "masv", "student id"])
    col_ten = find_column_index(header_row, ["họ và tên", "họ tên", "hoten", "full name"])
    col_lop = find_column_index(header_row, ["lớp", "class"])
    col_ngay = find_column_index(header_row, ["ngày sinh", "ngay sinh", "date of birth", "dob"])
    
    if col_ma is None or col_ten is None or col_lop is None or col_ngay is None:
        raise HTTPException(status_code=400, detail="File Excel thiếu cột bắt buộc: Mã sinh viên, Họ và tên, Lớp, Ngày sinh")
    
    report_rows = []
    success = 0
    failed = 0
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(row):
            continue
        ma_sv = str(row[col_ma]).strip() if col_ma < len(row) and row[col_ma] else ""
        ho_ten = str(row[col_ten]).strip() if col_ten < len(row) and row[col_ten] else ""
        lop = str(row[col_lop]).strip() if col_lop < len(row) and row[col_lop] else ""
        ngay_sinh_raw = row[col_ngay] if col_ngay < len(row) else None
        try:
            if not ma_sv:
                raise ValueError("Thiếu mã sinh viên")
            if not ho_ten:
                raise ValueError("Thiếu họ tên")
            if not lop:
                raise ValueError("Thiếu tên lớp")
            if not ngay_sinh_raw:
                raise ValueError("Thiếu ngày sinh")
            ngay_sinh_str = parse_date(ngay_sinh_raw)
            if not ngay_sinh_str:
                raise ValueError(f"Ngày sinh không hợp lệ: {ngay_sinh_raw} (cần dd/mm/yyyy)")
            # Mật khẩu mặc định: ngày sinh dạng ddmmyyyy (bỏ dấu /)
            mat_khau_raw = ngay_sinh_str.replace("/", "")
            existing = db.query(NguoiDung).filter(NguoiDung.ten_dang_nhap == ma_sv).first()
            if existing:
                raise ValueError("Mã sinh viên đã tồn tại")
            new_user = NguoiDung(
                ten_dang_nhap=ma_sv,
                mat_khau=hash_password(mat_khau_raw),
                ho_ten=ho_ten,
                vai_tro="hoc_sinh",
                kich_hoat=True,
                phai_doi_mat_khau=True
            )
            db.add(new_user)
            db.commit()
            success += 1
            report_rows.append([ma_sv, ho_ten, lop, "✅ Thành công", f"Mật khẩu: {mat_khau_raw}"])
        except Exception as e:
            failed += 1
            report_rows.append([ma_sv, ho_ten, lop, "❌ Thất bại", str(e)])
            db.rollback()
            continue
    
    wb_report = openpyxl.Workbook()
    ws_report = wb_report.active
    ws_report.title = "Báo cáo import"
    headers = ["Mã sinh viên", "Họ và tên", "Lớp", "Trạng thái", "Ghi chú"]
    for col, h in enumerate(headers, 1):
        cell = ws_report.cell(row=1, column=col)
        cell.value = h
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN
    for r_idx, row_data in enumerate(report_rows, start=2):
        for c_idx, val in enumerate(row_data, 1):
            cell = ws_report.cell(row=r_idx, column=c_idx)
            cell.value = val
            cell.font = Font(name="Times New Roman", size=11)
            cell.border = THIN
            if row_data[3] == "✅ Thành công":
                cell.fill = PatternFill("solid", fgColor="D5F5E3")
            else:
                cell.fill = PatternFill("solid", fgColor="FDECEA")
    output = io.BytesIO()
    wb_report.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=bao_cao_import_tai_khoan.xlsx"}
    )
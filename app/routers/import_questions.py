import io
import openpyxl
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form
from sqlalchemy.orm import Session
from app.core.database import get_db
from app.core.security import require_teacher
from app.models.models import CauHoi, LuaChon, ChuDe

router = APIRouter(prefix="/import", tags=["Import dữ liệu"])


@router.get("/template/excel", summary="Tải file Excel mẫu để nhập câu hỏi hàng loạt")
def tai_file_mau():
    """Trả về file Excel mẫu với hướng dẫn nhập câu hỏi."""
    from fastapi.responses import StreamingResponse
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Câu hỏi"

    HEADER_FILL = PatternFill("solid", fgColor="1B3A6B")
    HEADER_FONT = Font(bold=True, color="FFFFFF", name="Times New Roman", size=11)
    CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
    LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
    THIN   = Border(left=Side(style="thin"), right=Side(style="thin"),
                    top=Side(style="thin"),  bottom=Side(style="thin"))

    headers = ["chu_de_id", "noi_dung", "do_kho", "giai_thich", "dap_an_A", "dap_an_B", "dap_an_C", "dap_an_D", "dap_an_dung"]
    widths  = [12, 50, 14, 35, 22, 22, 22, 22, 14]

    for col, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=col)
        cell.value = h
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN
        ws.column_dimensions[cell.column_letter].width = w

    # Dòng hướng dẫn
    guide = ws.cell(row=2, column=1)
    guide.value = "← ID chủ đề"
    guide = ws.cell(row=2, column=3)
    guide.value = "de/trung_binh/kho"
    guide = ws.cell(row=2, column=9)
    guide.value = "A hoặc B hoặc C hoặc D"
    for col in range(1, 10):
        ws.cell(row=2, column=col).font = Font(italic=True, color="888888", name="Times New Roman", size=10)

    # Dòng mẫu
    sample = [6, "Kiểu dữ liệu nào dùng để lưu số nguyên trong Python?", "de",
              "int là viết tắt của integer - số nguyên",
              "int", "float", "str", "bool", "A"]
    for col, val in enumerate(sample, 1):
        cell = ws.cell(row=3, column=col)
        cell.value = val
        cell.font = Font(name="Times New Roman", size=11, color="1B3A6B")
        cell.alignment = LEFT
        cell.border = THIN
    ws.row_dimensions[3].height = 28

    # Sheet hướng dẫn
    ws2 = wb.create_sheet("Hướng dẫn")
    ws2["A1"] = "HƯỚNG DẪN NHẬP CÂU HỎI HÀNG LOẠT"
    ws2["A1"].font = Font(bold=True, size=14, color="1B3A6B", name="Times New Roman")
    huong_dan = [
        ("chu_de_id", "ID của chủ đề (xem trong hệ thống, ví dụ: 6 = Biến Python)"),
        ("noi_dung",  "Nội dung câu hỏi (bắt buộc, tối thiểu 5 ký tự)"),
        ("do_kho",    "Độ khó: de | trung_binh | kho (mặc định: trung_binh)"),
        ("giai_thich","Giải thích đáp án (không bắt buộc)"),
        ("dap_an_A",  "Nội dung đáp án A (bắt buộc)"),
        ("dap_an_B",  "Nội dung đáp án B (bắt buộc)"),
        ("dap_an_C",  "Nội dung đáp án C (không bắt buộc)"),
        ("dap_an_D",  "Nội dung đáp án D (không bắt buộc)"),
        ("dap_an_dung","Chữ cái đáp án đúng: A hoặc B hoặc C hoặc D (bắt buộc)"),
    ]
    for i, (col, desc) in enumerate(huong_dan, 3):
        ws2.cell(row=i, column=1).value = col
        ws2.cell(row=i, column=1).font = Font(bold=True, name="Times New Roman")
        ws2.cell(row=i, column=2).value = desc
        ws2.cell(row=i, column=2).font = Font(name="Times New Roman")
    ws2.column_dimensions["A"].width = 16
    ws2.column_dimensions["B"].width = 60

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    from fastapi.responses import StreamingResponse
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=mau_nhap_cau_hoi.xlsx"}
    )


@router.post("/questions/excel", summary="Import câu hỏi hàng loạt từ file Excel")
async def import_cau_hoi_excel(
    file: UploadFile = File(..., description="File Excel theo mẫu"),
    db: Session = Depends(get_db),
    current_user=Depends(require_teacher)
):
    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="Chỉ chấp nhận file .xlsx hoặc .xls")

    content = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content))
        ws = wb.active
    except Exception:
        raise HTTPException(status_code=400, detail="Không đọc được file Excel. Hãy dùng đúng file mẫu.")

    thanh_cong = []
    loi = []
    DO_KHO_VALID = {"de", "trung_binh", "kho"}
    DAP_AN_MAP = {"A": 0, "B": 1, "C": 2, "D": 3}

    for row_idx, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
        # Bỏ qua dòng trống
        if not any(row):
            continue

        try:
            chu_de_id   = int(row[0]) if row[0] else None
            noi_dung    = str(row[1]).strip() if row[1] else ""
            do_kho      = str(row[2]).strip().lower() if row[2] else "trung_binh"
            giai_thich  = str(row[3]).strip() if row[3] else None
            dap_an_A    = str(row[4]).strip() if row[4] else ""
            dap_an_B    = str(row[5]).strip() if row[5] else ""
            dap_an_C    = str(row[6]).strip() if row[6] else ""
            dap_an_D    = str(row[7]).strip() if row[7] else ""
            dap_an_dung = str(row[8]).strip().upper() if row[8] else ""

            # Validate
            if not chu_de_id:
                raise ValueError("Thiếu chu_de_id")
            if len(noi_dung) < 5:
                raise ValueError("Nội dung câu hỏi quá ngắn")
            if do_kho not in DO_KHO_VALID:
                do_kho = "trung_binh"
            if not dap_an_A or not dap_an_B:
                raise ValueError("Phải có ít nhất đáp án A và B")
            if dap_an_dung not in DAP_AN_MAP:
                raise ValueError(f"Đáp án đúng phải là A, B, C hoặc D")

            chu_de = db.query(ChuDe).filter(ChuDe.id == chu_de_id).first()
            if not chu_de:
                raise ValueError(f"Chủ đề ID={chu_de_id} không tồn tại")

            cau_hoi = CauHoi(
                noi_dung=noi_dung,
                loai_cau_hoi="trac_nghiem_mot_dap_an",
                do_kho=do_kho,
                chu_de_id=chu_de_id,
                nguoi_tao_id=current_user.id,
                giai_thich=giai_thich or None,
                an_hien=True,
            )
            db.add(cau_hoi)
            db.flush()

            cac_dap_an = [dap_an_A, dap_an_B, dap_an_C, dap_an_D]
            idx_dung = DAP_AN_MAP[dap_an_dung]
            for i, nd in enumerate(cac_dap_an):
                if not nd:
                    continue
                lc = LuaChon(
                    noi_dung=nd,
                    la_dap_an=(i == idx_dung),
                    cau_hoi_id=cau_hoi.id,
                    thu_tu=i,
                )
                db.add(lc)

            thanh_cong.append({"dong": row_idx, "noi_dung": noi_dung[:50]})

        except Exception as e:
            loi.append({"dong": row_idx, "loi": str(e)})
            db.rollback()
            continue

    if thanh_cong:
        db.commit()

    return {
        "tong_dong": row_idx - 2 if 'row_idx' in dir() else 0,
        "thanh_cong": len(thanh_cong),
        "that_bai": len(loi),
        "chi_tiet_loi": loi,
    }